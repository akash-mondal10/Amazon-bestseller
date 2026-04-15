from fastapi import FastAPI, APIRouter, HTTPException, Query
from fastapi.responses import StreamingResponse
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
import os
import logging
import io
from pathlib import Path
from pydantic import BaseModel
from typing import List, Optional
import sys

# Add amazon_analyzer to path
sys.path.insert(0, str(Path(__file__).parent / 'amazon_analyzer' / 'src'))

from data_loader import DataLoader
from preprocessing import DataPreprocessor
from analysis import DataAnalyzer
from ml_predictor import BookPredictor

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

# Initialize data once at startup
print("Loading and preprocessing Amazon Bestsellers dataset...")
loader = DataLoader()
df = loader.load_data()
preprocessor = DataPreprocessor(df)
preprocessor.check_duplicates()
preprocessor.remove_duplicates()
preprocessor.validate_data_types()
clean_df = preprocessor.get_clean_data()
analyzer = DataAnalyzer(clean_df)
print(f"Dataset ready: {len(clean_df)} books loaded")

# Initialize ML predictor
print("Training ML models...")
predictor = BookPredictor(clean_df)
print("ML models trained successfully")

# Create the main app
app = FastAPI(title="Amazon Bestseller Analyzer API")

# Create a router with the /api prefix
api_router = APIRouter(prefix="/api")


# Response Models
class SummaryStats(BaseModel):
    total_books: int
    avg_rating: float
    avg_price: float
    total_reviews: int
    genres: List[str]
    year_range: List[int]


class TopAuthors(BaseModel):
    authors: List[str]
    counts: List[int]


class GenreDistribution(BaseModel):
    genres: List[str]
    counts: List[int]
    percentages: List[float]


class YearTrends(BaseModel):
    years: List[int]
    book_counts: List[int]
    avg_ratings: List[float]
    avg_prices: List[float]
    avg_reviews: List[int]


class PriceRatingData(BaseModel):
    prices: List[float]
    ratings: List[float]
    names: List[str]
    genres: List[str]
    correlation: float


class CorrelationMatrix(BaseModel):
    columns: List[str]
    matrix: List[List[float]]


class FilterOptions(BaseModel):
    years: List[int]
    genres: List[str]


class PredictionRequest(BaseModel):
    price: float
    year: int
    genre: str


# API Endpoints
@api_router.get("/")
async def root():
    return {
        "message": "Amazon Bestseller Analyzer API",
        "version": "2.0",
        "total_books": len(clean_df)
    }


@api_router.get("/summary", response_model=SummaryStats)
async def get_summary():
    """Get summary statistics of the dataset."""
    return {
        "total_books": len(clean_df),
        "avg_rating": float(clean_df['User Rating'].mean()),
        "avg_price": float(clean_df['Price'].mean()),
        "total_reviews": int(clean_df['Reviews'].sum()),
        "genres": clean_df['Genre'].unique().tolist(),
        "year_range": [int(clean_df['Year'].min()), int(clean_df['Year'].max())]
    }


@api_router.get("/top-authors", response_model=TopAuthors)
async def get_top_authors(n: int = 10):
    """Get top N authors by book count."""
    return analyzer.get_top_authors(n)


@api_router.get("/genre-distribution", response_model=GenreDistribution)
async def get_genre_distribution():
    """Get distribution of books by genre."""
    return analyzer.get_genre_distribution()


@api_router.get("/year-trends", response_model=YearTrends)
async def get_year_trends():
    """Get year-wise trends."""
    return analyzer.get_year_wise_trends()


@api_router.get("/price-rating", response_model=PriceRatingData)
async def get_price_rating():
    """Get price vs rating data."""
    price_rating = analyzer.get_price_rating_analysis()
    price_rating['genres'] = clean_df['Genre'].tolist()
    return price_rating


@api_router.get("/correlation", response_model=CorrelationMatrix)
async def get_correlation():
    """Get correlation matrix."""
    return analyzer.get_correlation_matrix()


@api_router.get("/books")
async def get_books(
    year: Optional[int] = None,
    genre: Optional[str] = None,
    search: Optional[str] = None,
    limit: int = 100
):
    """Get filtered/searched books data."""
    filtered_df = analyzer.get_filtered_data(year=year, genre=genre)

    # Apply search filter
    if search and search.strip():
        query = search.strip().lower()
        mask = (
            filtered_df['Name'].str.lower().str.contains(query, na=False) |
            filtered_df['Author'].str.lower().str.contains(query, na=False)
        )
        filtered_df = filtered_df[mask]

    books = filtered_df.head(limit).to_dict('records')

    return {
        "total": len(filtered_df),
        "showing": len(books),
        "books": books
    }


@api_router.get("/search")
async def search_books(
    q: str = Query(..., min_length=1, description="Search query"),
    limit: int = 20
):
    """Search books by name or author."""
    query = q.strip().lower()
    mask = (
        clean_df['Name'].str.lower().str.contains(query, na=False) |
        clean_df['Author'].str.lower().str.contains(query, na=False)
    )
    results = clean_df[mask].head(limit)

    return {
        "query": q,
        "total": int(mask.sum()),
        "results": results.to_dict('records')
    }


@api_router.get("/insights")
async def get_insights():
    """Get generated insights from the data."""
    insights = analyzer.generate_insights()
    return {"insights": insights}


@api_router.get("/filter-options", response_model=FilterOptions)
async def get_filter_options():
    """Get available filter options."""
    return {
        "years": sorted(clean_df['Year'].unique().tolist()),
        "genres": sorted(clean_df['Genre'].unique().tolist())
    }


# ML Prediction Endpoints
@api_router.get("/ml/info")
async def get_ml_info():
    """Get ML model information and metrics."""
    return predictor.get_model_info()


@api_router.post("/ml/predict")
async def predict(request: PredictionRequest):
    """Predict bestseller status and rating for a book."""
    if request.genre not in ['Fiction', 'Non Fiction']:
        raise HTTPException(status_code=400, detail="Genre must be 'Fiction' or 'Non Fiction'")

    bestseller_pred = predictor.predict_bestseller(request.price, request.year, request.genre)
    rating_pred = predictor.predict_rating(request.price, request.year, request.genre)

    return {
        "input": {
            "price": request.price,
            "year": request.year,
            "genre": request.genre
        },
        "bestseller_prediction": bestseller_pred,
        "rating_prediction": rating_pred
    }


# Export Endpoints
@api_router.get("/export/excel")
async def export_excel(
    year: Optional[int] = None,
    genre: Optional[str] = None,
    search: Optional[str] = None
):
    """Export filtered data as Excel with analysis sheets."""
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.chart import BarChart, Reference, PieChart, LineChart
    import pandas as pd

    filtered_df = analyzer.get_filtered_data(year=year, genre=genre)
    if search and search.strip():
        query = search.strip().lower()
        mask = (
            filtered_df['Name'].str.lower().str.contains(query, na=False) |
            filtered_df['Author'].str.lower().str.contains(query, na=False)
        )
        filtered_df = filtered_df[mask]

    wb = openpyxl.Workbook()

    # Style constants
    header_font = Font(bold=True, size=12, color="FFFFFF")
    header_fill = PatternFill(start_color="EA580C", end_color="EA580C", fill_type="solid")
    title_font = Font(bold=True, size=16, color="030712")
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    # Sheet 1: Summary
    ws_summary = wb.active
    ws_summary.title = "Summary"
    ws_summary['A1'] = "Amazon Bestseller Analysis Report"
    ws_summary['A1'].font = title_font
    ws_summary.merge_cells('A1:D1')

    summary_data = [
        ("Total Books", len(filtered_df)),
        ("Average Rating", round(float(filtered_df['User Rating'].mean()), 2)),
        ("Average Price", f"${float(filtered_df['Price'].mean()):.2f}"),
        ("Total Reviews", int(filtered_df['Reviews'].sum())),
        ("Year Range", f"{int(filtered_df['Year'].min())} - {int(filtered_df['Year'].max())}"),
        ("Genres", ", ".join(filtered_df['Genre'].unique().tolist())),
    ]
    for i, (key, val) in enumerate(summary_data, start=3):
        ws_summary[f'A{i}'] = key
        ws_summary[f'A{i}'].font = Font(bold=True)
        ws_summary[f'B{i}'] = str(val)

    # Insights
    ws_summary[f'A{len(summary_data)+4}'] = "Key Insights"
    ws_summary[f'A{len(summary_data)+4}'].font = Font(bold=True, size=14)
    insights = analyzer.generate_insights()
    for i, insight in enumerate(insights, start=len(summary_data)+5):
        ws_summary[f'A{i}'] = f"{i-len(summary_data)-4}. {insight}"

    ws_summary.column_dimensions['A'].width = 50
    ws_summary.column_dimensions['B'].width = 25

    # Sheet 2: Books Data
    ws_data = wb.create_sheet("Books Data")
    headers = ['Name', 'Author', 'User Rating', 'Reviews', 'Price', 'Year', 'Genre']
    for col, header in enumerate(headers, 1):
        cell = ws_data.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border

    for row_idx, (_, row) in enumerate(filtered_df[headers].iterrows(), start=2):
        for col_idx, val in enumerate(row, 1):
            cell = ws_data.cell(row=row_idx, column=col_idx, value=val)
            cell.border = thin_border

    for col in ['A', 'B', 'C', 'D', 'E', 'F', 'G']:
        ws_data.column_dimensions[col].width = 35 if col in ['A', 'B'] else 15

    # Sheet 3: Top Authors
    ws_authors = wb.create_sheet("Top Authors")
    top = analyzer.get_top_authors(15)
    ws_authors['A1'] = "Author"
    ws_authors['B1'] = "Book Count"
    ws_authors['A1'].font = header_font
    ws_authors['A1'].fill = header_fill
    ws_authors['B1'].font = header_font
    ws_authors['B1'].fill = header_fill
    for i, (author, count) in enumerate(zip(top['authors'], top['counts']), 2):
        ws_authors[f'A{i}'] = author
        ws_authors[f'B{i}'] = count
    ws_authors.column_dimensions['A'].width = 35
    ws_authors.column_dimensions['B'].width = 15

    chart = BarChart()
    chart.title = "Top Authors by Book Count"
    chart.y_axis.title = "Number of Books"
    data_ref = Reference(ws_authors, min_col=2, min_row=1, max_row=len(top['authors'])+1)
    cats_ref = Reference(ws_authors, min_col=1, min_row=2, max_row=len(top['authors'])+1)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)
    chart.width = 20
    chart.height = 12
    ws_authors.add_chart(chart, "D2")

    # Sheet 4: Year Trends
    ws_trends = wb.create_sheet("Year Trends")
    trends = analyzer.get_year_wise_trends()
    ws_trends['A1'] = "Year"
    ws_trends['B1'] = "Book Count"
    ws_trends['C1'] = "Avg Rating"
    ws_trends['D1'] = "Avg Price"
    for col in ['A', 'B', 'C', 'D']:
        ws_trends[f'{col}1'].font = header_font
        ws_trends[f'{col}1'].fill = header_fill
    for i, (yr, cnt, rt, pr) in enumerate(zip(
        trends['years'], trends['book_counts'], trends['avg_ratings'], trends['avg_prices']
    ), 2):
        ws_trends[f'A{i}'] = yr
        ws_trends[f'B{i}'] = cnt
        ws_trends[f'C{i}'] = rt
        ws_trends[f'D{i}'] = pr

    line_chart = LineChart()
    line_chart.title = "Year-wise Book Count Trends"
    data_ref = Reference(ws_trends, min_col=2, min_row=1, max_row=len(trends['years'])+1)
    cats_ref = Reference(ws_trends, min_col=1, min_row=2, max_row=len(trends['years'])+1)
    line_chart.add_data(data_ref, titles_from_data=True)
    line_chart.set_categories(cats_ref)
    line_chart.width = 20
    line_chart.height = 12
    ws_trends.add_chart(line_chart, "F2")

    # Sheet 5: Correlation
    ws_corr = wb.create_sheet("Correlation")
    corr_data = analyzer.get_correlation_matrix()
    ws_corr['A1'] = ""
    for i, col in enumerate(corr_data['columns'], 2):
        ws_corr.cell(row=1, column=i, value=col).font = header_font
        ws_corr.cell(row=1, column=i).fill = header_fill
    for r, (row_label, row_vals) in enumerate(zip(corr_data['columns'], corr_data['matrix']), 2):
        ws_corr.cell(row=r, column=1, value=row_label).font = Font(bold=True)
        for c, val in enumerate(row_vals, 2):
            cell = ws_corr.cell(row=r, column=c, value=round(val, 3))
            if val > 0.3:
                cell.fill = PatternFill(start_color="FFB3B3", end_color="FFB3B3", fill_type="solid")
            elif val < -0.3:
                cell.fill = PatternFill(start_color="B3B3FF", end_color="B3B3FF", fill_type="solid")

    # Save to buffer
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=amazon_bestseller_analysis.xlsx"}
    )


@api_router.get("/export/pdf")
async def export_pdf(
    year: Optional[int] = None,
    genre: Optional[str] = None,
    search: Optional[str] = None
):
    """Export analysis report as PDF."""
    from reportlab.lib.pagesizes import letter, A4
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
    from reportlab.lib.units import inch
    import pandas as pd

    filtered_df = analyzer.get_filtered_data(year=year, genre=genre)
    if search and search.strip():
        query = search.strip().lower()
        mask = (
            filtered_df['Name'].str.lower().str.contains(query, na=False) |
            filtered_df['Author'].str.lower().str.contains(query, na=False)
        )
        filtered_df = filtered_df[mask]

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=0.75*inch, bottomMargin=0.75*inch)
    styles = getSampleStyleSheet()
    story = []

    # Custom styles
    title_style = ParagraphStyle('CustomTitle', parent=styles['Title'], fontSize=24, textColor=colors.HexColor('#EA580C'), spaceAfter=20)
    heading_style = ParagraphStyle('CustomHeading', parent=styles['Heading2'], fontSize=16, textColor=colors.HexColor('#030712'), spaceBefore=20, spaceAfter=10)
    body_style = ParagraphStyle('CustomBody', parent=styles['Normal'], fontSize=11, leading=16)
    insight_style = ParagraphStyle('InsightStyle', parent=styles['Normal'], fontSize=10, leading=14, leftIndent=20, spaceAfter=6)

    # Title
    story.append(Paragraph("Amazon Bestseller Analyzer", title_style))
    story.append(Paragraph("Comprehensive Analysis Report", styles['Heading3']))
    story.append(Spacer(1, 20))

    # Summary Statistics
    story.append(Paragraph("Summary Statistics", heading_style))
    summary_data = [
        ['Metric', 'Value'],
        ['Total Books', str(len(filtered_df))],
        ['Average Rating', f"{float(filtered_df['User Rating'].mean()):.2f}"],
        ['Average Price', f"${float(filtered_df['Price'].mean()):.2f}"],
        ['Total Reviews', f"{int(filtered_df['Reviews'].sum()):,}"],
        ['Year Range', f"{int(filtered_df['Year'].min())} - {int(filtered_df['Year'].max())}"],
        ['Fiction Books', str(len(filtered_df[filtered_df['Genre'] == 'Fiction']))],
        ['Non-Fiction Books', str(len(filtered_df[filtered_df['Genre'] == 'Non Fiction']))],
    ]
    t = Table(summary_data, colWidths=[200, 200])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#EA580C')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 10),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E5E7EB')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F9FAFB')]),
        ('TOPPADDING', (0, 0), (-1, -1), 8),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
    ]))
    story.append(t)
    story.append(Spacer(1, 20))

    # Key Insights
    story.append(Paragraph("Key Insights", heading_style))
    insights = analyzer.generate_insights()
    for i, insight in enumerate(insights, 1):
        story.append(Paragraph(f"<b>{i}.</b> {insight}", insight_style))
    story.append(Spacer(1, 20))

    # Top Authors
    story.append(Paragraph("Top 10 Authors", heading_style))
    top = analyzer.get_top_authors(10)
    author_data = [['Rank', 'Author', 'Books']]
    for i, (author, count) in enumerate(zip(top['authors'], top['counts']), 1):
        author_data.append([str(i), author, str(count)])
    t2 = Table(author_data, colWidths=[50, 250, 80])
    t2.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2563EB')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 10),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E5E7EB')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F9FAFB')]),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))
    story.append(t2)
    story.append(Spacer(1, 20))

    # Correlation Matrix
    story.append(Paragraph("Correlation Matrix", heading_style))
    corr = analyzer.get_correlation_matrix()
    corr_table = [[''] + corr['columns']]
    for i, row in enumerate(corr['matrix']):
        corr_table.append([corr['columns'][i]] + [f"{v:.3f}" for v in row])
    t3 = Table(corr_table, colWidths=[90]*5)
    t3.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#EA580C')),
        ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#EA580C')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('TEXTCOLOR', (0, 0), (0, -1), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('ALIGN', (1, 1), (-1, -1), 'CENTER'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E5E7EB')),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))
    story.append(t3)
    story.append(Spacer(1, 20))

    # ML Model Info
    story.append(Paragraph("ML Prediction Models", heading_style))
    ml_info = predictor.get_model_info()
    story.append(Paragraph(f"<b>Classification Model:</b> {ml_info['classification']['model']}", body_style))
    story.append(Paragraph(f"Task: {ml_info['classification']['task']}", insight_style))
    story.append(Paragraph(f"Accuracy: {ml_info['classification']['metrics']['accuracy']:.2%}", insight_style))
    story.append(Spacer(1, 10))
    story.append(Paragraph(f"<b>Regression Model:</b> {ml_info['regression']['model']}", body_style))
    story.append(Paragraph(f"Task: {ml_info['regression']['task']}", insight_style))
    story.append(Paragraph(f"RMSE: {ml_info['regression']['metrics']['rmse']:.4f}", insight_style))
    story.append(Paragraph(f"R2 Score: {ml_info['regression']['metrics']['r2_score']:.4f}", insight_style))
    story.append(Spacer(1, 20))

    # Page break before data table
    story.append(PageBreak())

    # Top Books Table
    story.append(Paragraph("Top 20 Books by Reviews", heading_style))
    top_books = filtered_df.nlargest(20, 'Reviews')[['Name', 'Author', 'User Rating', 'Reviews', 'Price', 'Genre']]
    book_headers = ['Name', 'Author', 'Rating', 'Reviews', 'Price', 'Genre']
    book_data = [book_headers]
    for _, row in top_books.iterrows():
        name = str(row['Name'])[:30] + '...' if len(str(row['Name'])) > 30 else str(row['Name'])
        book_data.append([
            name, str(row['Author'])[:20], str(row['User Rating']),
            f"{int(row['Reviews']):,}", f"${int(row['Price'])}", str(row['Genre'])
        ])
    t4 = Table(book_data, colWidths=[120, 85, 40, 65, 40, 65])
    t4.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#10B981')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E5E7EB')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F9FAFB')]),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ]))
    story.append(t4)

    # Footer
    story.append(Spacer(1, 30))
    story.append(Paragraph("Generated by Amazon Bestseller Analyzer | Python Data Analysis Project", 
                           ParagraphStyle('Footer', parent=styles['Normal'], fontSize=8, textColor=colors.grey)))

    doc.build(story)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": "attachment; filename=amazon_bestseller_report.pdf"}
    )


# Include the router in the main app
app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get('CORS_ORIGINS', '*').split(','),
    allow_methods=["*"],
    allow_headers=["*"],
)

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)
